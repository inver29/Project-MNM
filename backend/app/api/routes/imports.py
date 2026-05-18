from __future__ import annotations

from io import BytesIO

from fastapi import APIRouter, Depends, File, Form, HTTPException, Response, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import selectinload

from app.api.deps import CurrentAccount, DbSession, require_roles
from app.models import AccountRole, CatalogItem, CatalogSpace, PurchaseImportBatch, PurchaseImportLine
from app.schemas.contracts import PurchaseImportLineRead, PurchaseImportRead
from app.services.inventory import build_import_invoice_code, ensure_inventory_baseline, lock_catalog_items


router = APIRouter()

HEADER_ALIASES = {
    "item_id": {"item_id", "ma_san_pham"},
    "title": {"title", "ten_san_pham", "product_title"},
    "space_name": {"space_name", "ten_khong_gian", "category", "danh_muc"},
    "quantity": {"quantity", "so_luong", "so_luong_nhap"},
    "import_price": {"import_price", "gia_nhap"},
    "sale_price": {"sale_price", "gia_ban"},
    "summary_text": {"summary_text", "mo_ta"},
    "material_note": {"material_note", "chat_lieu"},
    "color_tone": {"color_tone", "mau_sac"},
    "dimension_note": {"dimension_note", "kich_thuoc"},
    "primary_image_url": {"primary_image_url", "anh_dai_dien", "image_url"},
    "is_published": {"is_published", "dang_kinh_doanh"},
}


def normalize_header(value: object) -> str:
    return str(value or "").strip().lower().replace(" ", "_")


def parse_bool(value: object) -> bool | None:
    if value is None or value == "":
        return None
    normalized = str(value).strip().lower()
    if normalized in {"1", "true", "yes", "co", "x"}:
        return True
    if normalized in {"0", "false", "no", "khong"}:
        return False
    return None


def parse_int(value: object, field_label: str, *, allow_empty: bool = False) -> int | None:
    if value is None or value == "":
        return None if allow_empty else 0
    try:
        return int(float(value))
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=f"Giá trị '{field_label}' trong file Excel không hợp lệ.") from exc


def resolve_headers(header_row: tuple[object, ...]) -> dict[str, int]:
    resolved: dict[str, int] = {}
    for index, value in enumerate(header_row):
        normalized = normalize_header(value)
        for target, aliases in HEADER_ALIASES.items():
            if normalized in aliases:
                resolved[target] = index
                break
    if "quantity" not in resolved:
        raise HTTPException(status_code=400, detail="File Excel phải có cột 'quantity' hoặc 'so_luong'.")
    if "item_id" not in resolved and not {"title", "space_name"}.issubset(resolved):
        raise HTTPException(
            status_code=400,
            detail="Mỗi dòng trong file Excel cần có 'item_id' hoặc đồng thời 'title' và 'space_name'.",
        )
    return resolved


def get_cell(row: tuple[object, ...], header_map: dict[str, int], key: str):
    index = header_map.get(key)
    if index is None or index >= len(row):
        return None
    return row[index]


def build_purchase_import_response(batch: PurchaseImportBatch) -> PurchaseImportRead:
    lines = list(batch.items)
    return PurchaseImportRead(
        batch_id=batch.batch_id,
        invoice_code=batch.invoice_code,
        imported_by_account_id=batch.imported_by_account_id,
        imported_by_display_name=batch.imported_by.display_name if batch.imported_by else None,
        note=batch.note,
        created_at=batch.created_at,
        updated_at=batch.updated_at,
        items=[
            PurchaseImportLineRead(
                import_line_id=line.import_line_id,
                item_id=line.item_id,
                item_title_snapshot=line.item_title_snapshot,
                imported_qty=line.imported_qty,
                remaining_qty=line.remaining_qty,
                import_unit_price=int(line.import_unit_price),
                sale_unit_price_snapshot=int(line.sale_unit_price_snapshot),
                note=line.note,
                created_at=line.created_at,
            )
            for line in lines
        ],
        total_import_qty=sum(line.imported_qty for line in lines),
        total_import_amount=sum(int(line.import_unit_price) * line.imported_qty for line in lines),
    )


def purchase_import_query():
    return select(PurchaseImportBatch).options(
        selectinload(PurchaseImportBatch.imported_by),
        selectinload(PurchaseImportBatch.items).selectinload(PurchaseImportLine.inventory_allocations),
    )


def get_batch_or_404(db: DbSession, batch_id: int, *, lock: bool = False) -> PurchaseImportBatch:
    statement = purchase_import_query().where(PurchaseImportBatch.batch_id == batch_id)
    if lock:
        statement = statement.with_for_update()
    batch = db.scalar(statement)
    if not batch:
        raise HTTPException(status_code=404, detail="Không tìm thấy phiếu nhập hàng.")
    return batch


@router.get(
    "/inventory/imports",
    response_model=list[PurchaseImportRead],
    dependencies=[Depends(require_roles(AccountRole.admin, AccountRole.nhan_vien))],
)
def list_purchase_imports(db: DbSession) -> list[PurchaseImportRead]:
    batches = list(db.scalars(purchase_import_query().order_by(PurchaseImportBatch.created_at.desc())).unique())
    return [build_purchase_import_response(batch) for batch in batches]


@router.get(
    "/inventory/imports/{batch_id}",
    response_model=PurchaseImportRead,
    dependencies=[Depends(require_roles(AccountRole.admin, AccountRole.nhan_vien))],
)
def get_purchase_import(batch_id: int, db: DbSession) -> PurchaseImportRead:
    return build_purchase_import_response(get_batch_or_404(db, batch_id))


@router.post(
    "/inventory/imports/excel",
    response_model=PurchaseImportRead,
    status_code=status.HTTP_201_CREATED,
    dependencies=[Depends(require_roles(AccountRole.admin, AccountRole.nhan_vien))],
)
async def import_inventory_from_excel(
    db: DbSession,
    current_account: CurrentAccount,
    file: UploadFile = File(...),
    invoice_code: str | None = Form(default=None),
    note: str | None = Form(default=None),
) -> PurchaseImportRead:
    file_name = file.filename or "inventory.xlsx"
    if not file_name.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        raise HTTPException(status_code=400, detail="Chỉ hỗ trợ file Excel định dạng .xlsx hoặc .xlsm.")

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=400, detail="File Excel tải lên đang rỗng.")

    try:
        workbook = load_workbook(filename=BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Không đọc được file Excel đã tải lên.") from exc

    worksheet = workbook.active
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        raise HTTPException(status_code=400, detail="File Excel không có dòng tiêu đề.")

    header_map = resolve_headers(header_row)
    row_payloads: list[dict[str, object]] = []

    for row_index, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in row):
            continue

        item_id = parse_int(get_cell(row, header_map, "item_id"), "item_id", allow_empty=True)
        title = str(get_cell(row, header_map, "title") or "").strip()
        space_name = str(get_cell(row, header_map, "space_name") or "").strip()
        quantity = parse_int(get_cell(row, header_map, "quantity"), "quantity")
        import_price = parse_int(get_cell(row, header_map, "import_price"), "import_price", allow_empty=True) or 0
        sale_price = parse_int(get_cell(row, header_map, "sale_price"), "sale_price", allow_empty=True)

        if not quantity or quantity <= 0:
            raise HTTPException(status_code=400, detail=f"Dòng Excel {row_index} phải có số lượng nhập lớn hơn 0.")
        if not item_id and (not title or not space_name):
            raise HTTPException(
                status_code=400,
                detail=f"Dòng Excel {row_index} cần có 'item_id' hoặc đồng thời 'title' và 'space_name'.",
            )

        row_payloads.append(
            {
                "item_id": item_id,
                "title": title,
                "space_name": space_name,
                "quantity": quantity,
                "import_price": import_price,
                "sale_price": sale_price,
                "summary_text": str(get_cell(row, header_map, "summary_text") or "").strip(),
                "material_note": str(get_cell(row, header_map, "material_note") or "").strip(),
                "color_tone": str(get_cell(row, header_map, "color_tone") or "").strip(),
                "dimension_note": str(get_cell(row, header_map, "dimension_note") or "").strip(),
                "primary_image_url": str(get_cell(row, header_map, "primary_image_url") or "").strip() or None,
                "is_published": parse_bool(get_cell(row, header_map, "is_published")),
                "row_index": row_index,
            }
        )

    if not row_payloads:
        raise HTTPException(status_code=400, detail="File Excel không có dòng dữ liệu hợp lệ.")

    try:
        ensure_inventory_baseline(db)

        normalized_invoice_code = (invoice_code or "").strip() or build_import_invoice_code()
        existing_code = db.scalar(
            select(PurchaseImportBatch).where(PurchaseImportBatch.invoice_code == normalized_invoice_code)
        )
        if existing_code:
            raise HTTPException(status_code=400, detail="Mã hóa đơn nhập đã tồn tại.")

        batch = PurchaseImportBatch(
            invoice_code=normalized_invoice_code,
            imported_by_account_id=current_account.account_id,
            note=(note or "").strip() or f"Nhập hàng từ file Excel {file_name}",
        )
        db.add(batch)
        db.flush()

        locked_item_ids = sorted({int(payload["item_id"]) for payload in row_payloads if payload["item_id"]})
        locked_items = lock_catalog_items(db, locked_item_ids) if locked_item_ids else {}

        spaces_by_name = {
            space.space_name.lower(): space
            for space in db.scalars(select(CatalogSpace).order_by(CatalogSpace.space_name.asc()))
        }

        for payload in row_payloads:
            item = None
            item_id = payload["item_id"]
            if item_id:
                item = locked_items.get(int(item_id))
                if not item:
                    raise HTTPException(
                        status_code=400,
                        detail=f"Dòng Excel {payload['row_index']}: không tìm thấy sản phẩm có item_id {item_id}.",
                    )
            else:
                matched_space = spaces_by_name.get(str(payload["space_name"]).lower())
                if not matched_space:
                    raise HTTPException(
                        status_code=400,
                        detail=f"Dòng Excel {payload['row_index']}: không tìm thấy danh mục '{payload['space_name']}'.",
                    )
                item = db.scalar(
                    select(CatalogItem)
                    .where(
                        CatalogItem.space_id == matched_space.space_id,
                        func.lower(CatalogItem.title) == str(payload["title"]).lower(),
                    )
                    .with_for_update()
                )
                if not item:
                    item = CatalogItem(
                        space_id=matched_space.space_id,
                        title=str(payload["title"]).strip(),
                        summary_text=str(payload["summary_text"]).strip(),
                        material_note=str(payload["material_note"]).strip(),
                        color_tone=str(payload["color_tone"]).strip(),
                        dimension_note=str(payload["dimension_note"]).strip(),
                        unit_price=int(payload["sale_price"] or payload["import_price"] or 0),
                        on_hand_qty=0,
                        primary_image_url=payload["primary_image_url"],
                        is_published=bool(payload["is_published"]) if payload["is_published"] is not None else True,
                    )
                    db.add(item)
                    db.flush()

            if payload["summary_text"]:
                item.summary_text = str(payload["summary_text"]).strip()
            if payload["material_note"]:
                item.material_note = str(payload["material_note"]).strip()
            if payload["color_tone"]:
                item.color_tone = str(payload["color_tone"]).strip()
            if payload["dimension_note"]:
                item.dimension_note = str(payload["dimension_note"]).strip()
            if payload["primary_image_url"]:
                item.primary_image_url = str(payload["primary_image_url"])
            if payload["is_published"] is not None:
                item.is_published = bool(payload["is_published"])
            if payload["sale_price"] is not None:
                item.unit_price = int(payload["sale_price"])

            item.on_hand_qty += int(payload["quantity"])

            db.add(
                PurchaseImportLine(
                    batch_id=batch.batch_id,
                    item_id=item.item_id,
                    item_title_snapshot=item.title,
                    imported_qty=int(payload["quantity"]),
                    remaining_qty=int(payload["quantity"]),
                    import_unit_price=int(payload["import_price"] or 0),
                    sale_unit_price_snapshot=int(item.unit_price),
                    note=f"Nhập từ dòng Excel {payload['row_index']}",
                )
            )

        db.commit()
    except Exception:
        db.rollback()
        raise

    return build_purchase_import_response(get_batch_or_404(db, batch.batch_id))


@router.delete(
    "/inventory/imports/{batch_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    response_model=None,
    response_class=Response,
    dependencies=[Depends(require_roles(AccountRole.admin, AccountRole.nhan_vien))],
)
def delete_purchase_import(batch_id: int, db: DbSession) -> None:
    batch = get_batch_or_404(db, batch_id, lock=True)

    try:
        affected_item_ids = sorted({line.item_id for line in batch.items})
        locked_items = lock_catalog_items(db, affected_item_ids) if affected_item_ids else {}

        for line in batch.items:
            if line.remaining_qty != line.imported_qty:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        f"Không thể xóa phiếu nhập {batch.invoice_code} vì sản phẩm '{line.item_title_snapshot}' "
                        "đã được sử dụng một phần."
                    ),
                )
            if line.inventory_allocations:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        f"Không thể xóa phiếu nhập {batch.invoice_code} vì sản phẩm '{line.item_title_snapshot}' "
                        "đã từng được dùng để xử lý đơn hàng."
                    ),
                )

            item = locked_items.get(line.item_id)
            if not item or item.on_hand_qty < line.imported_qty:
                raise HTTPException(
                    status_code=400,
                    detail=f"Không thể hoàn tác tồn kho cho sản phẩm '{line.item_title_snapshot}'.",
                )
            item.on_hand_qty -= line.imported_qty

        db.delete(batch)
        db.commit()
    except Exception:
        db.rollback()
        raise
